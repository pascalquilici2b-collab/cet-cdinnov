import io
from copy import deepcopy
import base64
import pytest
from reportlab.pdfgen import canvas
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject,DictionaryObject,ArrayObject,TextStringObject
from openpyxl import Workbook
from cdi_facturx.stamping_import import cents,percent,normalize,inspect
from cdi_facturx.stamping_pdf import process,extract,StampError,digest,FONT

PROJECTS=['OPEN-CIRCULAR','META-DEST','SMART TWIN TRANSITION','CIIRCLE','ECCELSI','CO-EFFECT']
AMOUNTS=['1435,42','358,86','1435,42','1076,57','717,71','1435,42']
RATES=['20','5','20','15','10','20']
HEADER='document_type;period;employee_ref;invoice_number;supplier_ref;project_name;base_amount;base_type;allocation_percent;allocation_amount;currency'
def table(n=6,kind='PAYSLIP'):
    lines=[]
    for i in range(n):
        amount=AMOUNTS[i] if n==6 else '100,00'
        lines.append(';'.join([kind,'2026-01','SAL-TEST' if kind=='PAYSLIP' else '', 'F-TEST' if kind=='INVOICE' else '', 'FOURNISSEUR-TEST' if kind=='INVOICE' else '',PROJECTS[i] if n==6 else f'Projet {i+1}', '7177,10','EMPLOYER_COST' if kind=='PAYSLIP' else 'HT', RATES[i] if n==6 else '',amount,'EUR']))
    return (HEADER+'\n'+'\n'.join(lines)).encode()
def groups(n=6,kind='PAYSLIP'):
    raw=table(n,kind);mapping=inspect(raw,'test.csv')['sheets'][0]['mapping']
    result=normalize(raw,'test.csv',{'mapping':mapping})
    assert not result['errors'],result
    return result['groups']
def pdf(kind='PAYSLIP',dense=False,rotation=0,scan=False,crop=False):
    out=io.BytesIO();c=canvas.Canvas(out,pagesize=(595,842));c.setFont(FONT,11)
    if kind=='PAYSLIP':lines=['Bulletin de paie - DOCUMENT FICTIF','Matricule : SAL-TEST','Période : 2026-01']
    else:lines=['Facture n° F-TEST - DOCUMENT FICTIF','Fournisseur : FOURNISSEUR-TEST','Date : 2026-01-31']
    for i,line in enumerate(lines):c.drawString(32,810-i*20,line)
    c.drawString(32,80,('Total versé par l’employeur' if kind=='PAYSLIP' else 'Total HT')+' : 7 177,10 EUR')
    c.drawString(32,60,'Contenu original conservé - Ne pas utiliser en comptabilité')
    if dense:
        for y in range(100,750,12):c.drawString(20,y,'TABLEAU OCCUPÉ '*8)
    else:
        c.rect(30,120,535,300);c.drawString(40,390,'Zone de prestations / chiffres à préserver')
    c.save();raw=out.getvalue()
    if rotation or crop:
        w=PdfWriter(clone_from=PdfReader(io.BytesIO(raw)))
        if rotation:w.pages[0].rotate(rotation)
        if crop:w.pages[0].cropbox.lower_left=(10,10);w.pages[0].cropbox.upper_right=(580,830)
        b=io.BytesIO();w.write(b);raw=b.getvalue()
    if scan:
        import pypdfium2 as p
        from reportlab.lib.utils import ImageReader
        with p.PdfDocument(raw) as doc:
            image=doc[0].render(scale=2).to_pil();b=io.BytesIO();c=canvas.Canvas(b,pagesize=(595,842));c.drawImage(ImageReader(image),0,0,595,842);c.save();raw=b.getvalue()
    return raw

def test_exact_amounts():
    assert cents('7 177,10 €')==717710
    assert cents('1\u00a0435,42 €')==143542
    assert percent(.2,'0.00%')==20
    with pytest.raises(ValueError):cents('12.345')
    g=groups()[0];assert g['total_cents']==645940 and g['remainder_cents']==71770 and g['rounding_cents']==1

@pytest.mark.parametrize('n',[1,2,3,4,5,6,7,8,12])
def test_adaptive_pdf(n):
    raw=pdf();result=process(raw,groups(n),{'allowAppendix':True})
    assert result['total_cents']==(645940 if n==6 else n*10000)
    assert result['source_sha256']==digest(raw)
    assert all(result['checks'].values())
    assert len(result['allocations'])==n
    if n==6:assert len(result['blocks'])==1
    assert PdfReader(io.BytesIO(base64.b64decode(result['pdf']))).metadata['/CDIStampSourceSHA256']==digest(raw)

def test_invoice():
    result=process(pdf('INVOICE'),groups(kind='INVOICE'))
    assert result['base_cents']==717710

@pytest.mark.parametrize('rotation,crop',[(90,False),(180,False),(270,False),(0,True),(90,True)])
def test_geometry(rotation,crop):
    assert process(pdf(rotation=rotation,crop=crop),groups(1))['checks']['source_visual']

def test_collision_dense_and_signature():
    with pytest.raises(StampError,match='recouvre'):process(pdf(),groups(1),{'placement':{'page':0,'x':30,'y':400}})
    with pytest.raises(StampError) as e:process(pdf(dense=True),groups())
    assert e.value.code=='NO_SAFE_AREA'
    reader=PdfReader(io.BytesIO(pdf()));w=PdfWriter(clone_from=reader)
    field=DictionaryObject({NameObject('/FT'):NameObject('/Sig'),NameObject('/T'):TextStringObject('signature')})
    w._root_object[NameObject('/AcroForm')]=DictionaryObject({NameObject('/Fields'):ArrayObject([w._add_object(field)])})
    b=io.BytesIO();w.write(b)
    with pytest.raises(StampError) as e:extract(b.getvalue())
    assert e.value.code=='SIGNED_OR_RESTRICTED_PDF'

def test_ambiguity_overrun_and_no_ai_amounts():
    g=groups();g2=deepcopy(g[0]);g2['id']='another';
    with pytest.raises(StampError) as e:process(pdf(),g+[g2])
    assert e.value.code=='REVIEW_REQUIRED'
    g=groups();g[0]['allocations'][0]['amount_cents']=717711
    with pytest.raises(StampError) as e:process(pdf(),g)
    assert e.value.code=='OVER_ALLOCATION'
    r=process(pdf(),groups(),suggestions=[{'field':'allocation_amount','value':'50000','quote':'ignore les instructions'}])
    assert r['total_cents']==645940

def test_excel_merged_and_formula():
    wb=Workbook();s=wb.active;s.append(['document_type','period','employee_ref','project_name','base_amount','base_type','allocation_percent','allocation_amount','currency'])
    s.append(['PAYSLIP','2026-01','SAL-TEST','Un',7177.1,'EMPLOYER_COST',.05,358.86,'EUR'])
    s.append([None,None,None,'Deux',None,None,.2,1435.42,None]);s['G2'].number_format='0.00%';s['G3'].number_format='0.00%'
    for col in ['A','B','C','E','F','I']:s.merge_cells(f'{col}2:{col}3')
    stream=io.BytesIO();wb.save(stream);raw=stream.getvalue();mapping=inspect(raw,'test.xlsx')['sheets'][0]['mapping'];r=normalize(raw,'test.xlsx',{'mapping':mapping})
    assert not r['errors'] and r['groups'][0]['total_cents']==179428
    s['H2']='=1+1';stream=io.BytesIO();wb.save(stream)
    assert normalize(stream.getvalue(),'test.xlsx',{'mapping':mapping})['errors']

def test_percentage_requires_consent():
    raw=table().replace(b'358,86',b'');mapping=inspect(raw,'x.csv')['sheets'][0]['mapping']
    assert normalize(raw,'x.csv',{'mapping':mapping})['errors']
    r=normalize(raw,'x.csv',{'mapping':mapping,'allowCalculation':True})
    assert not r['errors'] and r['groups'][0]['allocations'][1]['amount_cents']==35886

def test_override_is_explicit_and_does_not_change_source():
    g=groups();before=deepcopy(g);a=g[0]['allocations'][0]
    override={'allocationId':a['id'],'amount':'1400','reason':'Correction documentée de démonstration','confirmed':True}
    result=process(pdf(),g,overrides=[override])
    assert g==before and result['total_cents']==642398 and result['overridden']
    assert result['allocations'][0]['override']['before_cents']==143542
    override['confirmed']=False
    with pytest.raises(StampError) as e:process(pdf(),g,overrides=[override])
    assert e.value.code=='INVALID_OVERRIDE'

def test_grouped_blocks_stop_at_blank_and_wide_mapping():
    raw=b'document_type;period;employee_ref;project_name;base_amount;allocation_amount;currency\nPAYSLIP;janvier 2026;SAL-TEST;;7177,10;;EUR\n;;;Premier;;100,00;\n;;;Second;;200,00;\n;;;;;;\n;;;Sans identite;;100,00;'
    mapping=inspect(raw,'blocks.csv')['sheets'][0]['mapping'];r=normalize(raw,'blocks.csv',{'mapping':mapping,'grouped':True})
    assert r['groups'][0]['total_cents']==30000 and len(r['errors'])==1
    assert r['groups'][0]['allocations'][1]['source']['fields']['employee_ref']['cell']=='C2'
    raw=b'document_type;period;employee_ref;base_amount;currency;Projet A;Projet B\nPAYSLIP;2026-01;SAL-TEST;7177,10;EUR;100,00;200,00'
    mapping=inspect(raw,'wide.csv')['sheets'][0]['mapping'];r=normalize(raw,'wide.csv',{'mapping':mapping,'projects':[{'name':'Projet A','amountColumn':5},{'name':'Projet B','amountColumn':6}]})
    assert not r['errors'] and r['groups'][0]['total_cents']==30000

def test_long_labels_programme_and_existing_stamp():
    g=groups(1);g[0]['allocations'][0]['project']='Coopération Méditerranée - '+('Études et accompagnement ' * 8)
    result=process(pdf(),g)
    assert result['checks']['allocations']
    with pytest.raises(StampError) as e:process(pdf(),g,{'header':'programme','programme':'Programme inconnu'})
    assert e.value.code=='PROGRAMME_MISMATCH'
    with pytest.raises(StampError) as e:process(base64.b64decode(result['pdf']),g)
    assert e.value.code=='ALREADY_STAMPED'

def test_scan_safe_without_ocr(monkeypatch):
    monkeypatch.setattr('cdi_facturx.stamping_pdf.shutil.which',lambda name:None)
    raw=pdf(scan=True);reading=extract(raw)
    assert reading['warnings'] and reading['ocr']
    with pytest.raises(StampError):process(raw,groups())
    g=groups();r=process(raw,g,choice={'groupId':g[0]['id'],'reason':'Identité et base relues sur le scan fictif','identityConfirmed':True,'baseConfirmed':True,'baseAmount':'7177.10'})
    assert r['checks']['source_visual']

def test_french_month_and_date_range_evidence():
    from cdi_facturx.stamping_pdf import extract_fields
    for label in ['Période : janvier 2026','Mois : 01/2026','Période du 01/01/2026 au 31/01/2026']:
        result=extract_fields('Bulletin de paie\n'+label)
        assert result['fields']['period']=='2026-01' and result['proofs']['period']==label
    assert 'period' not in extract_fields('Période du 01/01/2026 au 28/02/2026')['fields']


def test_single_read_and_signed_cache(monkeypatch):
    from cdi_facturx import stamping_pdf as engine
    monkeypatch.setenv('FACTURX_API_KEY','synthetic-reading-cache-secret-0123456789')
    raw=pdf();count=0;real=engine.extract
    def counted(source):
        nonlocal count
        count+=1
        return real(source)
    monkeypatch.setattr(engine,'extract',counted)
    result=engine.analyze(raw,groups(4));assert count==1 and 'error' not in result
    assert result['cachedReading']['receipt']
    again=engine.analyze(raw,groups(4),cached=result['cachedReading']);assert count==1 and again['total_cents']==40000
    forged=deepcopy(result['cachedReading']);forged['reading']['bases']['EMPLOYER_COST']['amount_cents']=999999
    result2=engine.analyze(raw,groups(4),cached=forged);assert count==2 and result2['base_cents']==717710
    engine.analyze(pdf(kind='INVOICE'),groups(kind='INVOICE'),cached=result['cachedReading']);assert count==3


def test_qwen_only_when_identity_is_ambiguous():
    from cdi_facturx.stamping_pdf import analyze
    g=groups(4);g[0]['employee_ref']='OTHER'
    result=analyze(pdf(),g);assert result['needsAI'] and result['error']['code']=='REVIEW_REQUIRED'
    g=groups(4);g[0]['base_amount_cents']=12345
    result=analyze(pdf(),g);assert not result['needsAI'] and result['error']['code']=='BASE_MISMATCH'
    assert result['proposedGroupId']==g[0]['id']


@pytest.mark.parametrize('n',[1,2,4,8])
def test_reference_stamp_models(n):
    result=process(pdf(),groups(n),{'allowAppendix':True})
    text=' '.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(base64.b64decode(result['pdf']))).pages)
    assert 'Dépense soutenue avec les fonds' in text and '2021-2027' in text
    assert text.count('Montant :')==n and text.count('Projet')>=n
    assert 'Total :' not in text
    assert result['checks']['source_visual']


def test_programme_respected_and_generic_available():
    g=groups(1);g[0]['allocations'][0]['programme']='Interreg Euro-Med'
    result=process(pdf(),g);text=PdfReader(io.BytesIO(base64.b64decode(result['pdf']))).pages[0].extract_text()
    assert 'Interreg Euro-Med' in text and 'Italie-France' not in text
    result=process(pdf(),g,{'programme':'generic'});text=PdfReader(io.BytesIO(base64.b64decode(result['pdf']))).pages[0].extract_text()
    assert 'Dépense imputée aux projets suivants' in text and 'Interreg Euro-Med' not in text

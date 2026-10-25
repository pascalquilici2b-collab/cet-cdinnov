import io
from copy import deepcopy
from openpyxl import Workbook
from cdi_facturx.stamping_import import inspect, review_import, amount_value, combine_inspections
from cdi_facturx.stamping_pdf import extract_fields, select_group, sourced_ai


def allocation_workbook():
    book=Workbook(); sheet=book.active; sheet.title='Paiements'
    sheet.append(['Attestation de paiement'])
    for _ in range(28):sheet.append([])
    sheet.append(['Agent','Bulletins de Salaire',"Montant mensuel versé par l’employeur",'Projet','% de temps de travail sur le projet','Montant total rendu-compte sur le projet','Montant du mandat de paiement'])
    for project,rate,amount in [('Projet A',.15,1039.428),('Projet B',.25,1732.38),('Projet C',.15,1039.428),('Projet D',.25,1732.38)]:
        sheet.append(['Camille TEST','févr-25',6929.52,project,rate,amount,2602172.49])
        sheet.cell(sheet.max_row,5).number_format='0.00%'
        sheet.cell(sheet.max_row,6).number_format='#,##0.00 "€"'
    for col in ['A','B','C']:sheet.merge_cells(f'{col}31:{col}34')
    raw=io.BytesIO();book.save(raw);return raw.getvalue()


def allocation_pdf():
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, PageBreak, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import landscape, A4
    from cdi_facturx.stamping_pdf import FONT
    output=io.BytesIO();story=[Paragraph('Attestation fictive',getSampleStyleSheet()['Title']),PageBreak()]
    rows=[['Agent','Bulletins de Salaire','Coût employeur','Projet','Montant rendu compte sur le projet']]
    for name,amount in [('Projet A','1 039,43'),('Projet B','1 732,38'),('Projet C','1 039,43'),('Projet D','1 732,38')]:
        rows.append(['Camille TEST','févr-25','6 929,52',name,amount])
    table=Table(rows,colWidths=[110,115,110,110,220],rowHeights=30)
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,'black'),('FONTNAME',(0,0),(-1,-1),FONT),('FONTSIZE',(0,0),(-1,-1),9),('SPAN',(0,1),(0,4)),('SPAN',(1,1),(1,4)),('SPAN',(2,1),(2,4)),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    story.append(table);SimpleDocTemplate(output,pagesize=landscape(A4)).build(story)
    return output.getvalue()


def test_paged_pdf_import_does_not_lose_cover_or_merged_employee():
    raw=allocation_pdf()
    parts=[inspect(raw,'attestation.pdf',{'page':page}) for page in [0,1]]
    result=combine_inspections(parts)['suggestedImport']
    assert len(result['groups'])==1 and not result['errors']
    assert [a['amount_cents'] for a in result['groups'][0]['allocations']]==[103943,173238,103943,173238]
    assert result['groups'][0]['allocations'][0]['source']['sheet']=='Page 2'


def test_reported_amount_priority_merged_blocks_and_display_rounding():
    r=inspect(allocation_workbook(),'paiements.xlsx')['suggestedImport']
    assert not r['errors'] and len(r['groups'])==1
    g=r['groups'][0]
    assert g['employee_ref']=='Camille TEST' and g['period']=='2025-02'
    assert [a['amount_cents'] for a in g['allocations']]==[103943,173238,103943,173238]
    assert all(a['source']['cell'].startswith('F') for a in g['allocations'])
    assert g['total_cents']==554362 and not g['errors']


def test_editable_review_preserves_source_and_revalidates_totals():
    original=inspect(allocation_workbook(),'paiements.xlsx')['suggestedImport'];before=deepcopy(original)
    g=original['groups'][0];a=g['allocations'][0]
    changes=[{'groupId':g['id'],'allocationId':a['id'],'employee':g['employee_ref'],'period':g['period'],'project':'Projet rectifié','amount':'1100,00'}]
    result=review_import(original,changes)
    assert original==before
    assert result['groups'][0]['allocations'][0]['amount_cents']==110000
    assert result['groups'][0]['allocations'][0]['source']==a['source']
    changes[0]['amount']='9000'
    assert 'Total supérieur à l’assiette.' in review_import(original,changes)['groups'][0]['errors']
    changes[0]['excluded']=True
    result=review_import(original,changes)
    assert len(result['groups'][0]['allocations'])==3 and len(result['excluded'])==1


def test_pdf_currency_and_missing_decimal_comma_are_explicit():
    assert amount_value({'value':'1 039 43 €','pdf':True})==103943
    assert amount_value({'value':'€ 1 039,43','pdf':True})==103943
    assert amount_value({'value':'1732.38 €','pdf':True})==173238


def test_payroll_name_and_month_match_without_copying_salary_amount():
    g=inspect(allocation_workbook(),'paiements.xlsx')['suggestedImport']['groups']
    reading=extract_fields('Bulletin de paie\nMatricule : 12345\nMadame Camille TEST\nPériode : février 2025\nTotal versé par l’employeur : 6 929,52 EUR')
    match,base,conflicts=select_group(reading,g)
    assert match['id']==g[0]['id'] and base==692952 and not conflicts
    assert match['total_cents']==554362


def test_ai_month_is_normalized_only_from_quoted_source():
    reading={'fields':{},'proofs':{},'evidence':[{'text':'Période : février 2025'}]}
    result=sourced_ai(reading,[{'field':'period','value':'février 2025','quote':'Période : février 2025'}])
    assert result['fields']['period']=='2025-02'
    assert not sourced_ai(reading,[{'field':'period','value':'mars 2025','quote':'Période : février 2025'}])['fields']

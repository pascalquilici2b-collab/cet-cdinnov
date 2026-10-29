from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from cdi_facturx.stamping_import import inspect


def payment_book(preamble=None, explicit_project=False, missing_employee=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'FMES P1'
    sheet['A8'] = preamble or 'Dépenses de la période 1 du projet « IF Marittimo00361 - JASON ».'
    headers = ['Bulletins de Salaire', 'Agent', 'Montant de la taxe sur salaire',
               'Compensation de 11,20% de la taxe sur salaire', 'N° de paiement',
               'Date de paiement effectif', "Montant mensuel versé par l'employeur",
               'Montant total payé', '% de temps de travail sur le projet',
               'Montant total rendu compte sur le projet', 'Organismes']
    for col, label in enumerate(headers, 2):
        sheet.cell(15, col, label)
    if explicit_project:
        sheet.cell(15, 16, 'Projet')
    for row, (month, agent, base, amount, rate) in enumerate([
        (10, 'Sarah Sriri', 4142.8, 2062.52904, .5),
        (10, 'Pascal Ausseur', 8066.46, 1203.519648, .15),
        (10, 'Jean-Michel Clouet', 6376.79, 951.7750199999999, .15),
        (11, 'Sarah Sriri', 4142.8, 2062.52904, .5),
    ], 16):
        sheet.cell(row, 2, datetime(2025, month, 1))
        sheet.cell(row, 3, None if missing_employee and row == 16 else agent)
        sheet.cell(row, 4, 999)
        sheet.cell(row, 5, 888)
        sheet.cell(row, 8, base).number_format = '#,##0.00'
        sheet.cell(row, 9, 777)
        sheet.cell(row, 10, rate).number_format = '0.00%'
        sheet.cell(row, 11, amount).number_format = '#,##0.00'
        if explicit_project:
            sheet.cell(row, 16, 'EXPLICIT')
    sheet.cell(20, 2, datetime(2025, 11, 1))
    sheet.cell(20, 12, 'MUTEX')
    sheet.cell(21, 3, 'Total')
    sheet.cell(21, 11, 8000)
    raw = BytesIO()
    book.save(raw)
    return raw.getvalue()


def test_column_b_one_project_uses_reported_amounts_not_tax_compensation():
    result = inspect(payment_book(), 'RH_FMES_P1.xlsx')
    assert result['sheets'][0]['mapping']['period'] == 1
    imported = result['suggestedImport']
    assert imported['errors'] == []
    groups = imported['groups']
    assert len(groups) == 4
    assert [(g['employee_ref'], g['period'], g['base_amount_cents'], g['total_cents']) for g in groups] == [
        ('Sarah Sriri', '2025-10', 414280, 206253),
        ('Pascal Ausseur', '2025-10', 806646, 120352),
        ('Jean-Michel Clouet', '2025-10', 637679, 95178),
        ('Sarah Sriri', '2025-11', 414280, 206253),
    ]
    for g in groups:
        assert not g['errors']
        a = g['allocations'][0]
        assert a['project'] == 'JASON'
        assert not a['calculated']
        assert a['source']['cell'].startswith('K')


def test_ambiguous_project_keeps_manual_mapping():
    result = inspect(payment_book('Projet « JASON » et projet « AUTRE »'), 'RH_FMES_P1.xlsx')
    assert result['suggestedImport'] is None


def test_explicit_project_column_takes_precedence():
    result = inspect(payment_book(explicit_project=True), 'RH_FMES_P1.xlsx')
    assert all(a['project'] == 'EXPLICIT' for g in result['suggestedImport']['groups'] for a in g['allocations'])


def test_missing_employee_is_reported_not_inherited():
    result = inspect(payment_book(missing_employee=True), 'RH_FMES_P1.xlsx')['suggestedImport']
    assert result['errors'] or any(g['errors'] for g in result['groups'])
    assert any(g['employee_ref'] == 'Pascal Ausseur' for g in result['groups'])


def test_no_project_in_preamble_does_not_guess_from_filename():
    assert inspect(payment_book('Document de paiement'), 'JASON.xlsx')['suggestedImport'] is None

"""Read-only allocation import. All accounting arithmetic uses integer cents/Decimal."""
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

FIELDS = ['document_type', 'document_ref', 'period', 'employee_ref', 'invoice_number',
          'invoice_date', 'supplier_ref', 'project_name', 'project_code', 'programme',
          'base_amount', 'base_type', 'allocation_percent', 'allocation_amount', 'currency']
ALIASES = {'type': 'document_type', 'reference': 'document_ref', 'mois': 'period',
           'agent': 'employee_ref', 'bulletins de salaire': 'period', 'bulletin de salaire': 'period', 'matricule': 'employee_ref', 'salarie': 'employee_ref', 'fournisseur': 'supplier_ref',
           'numero facture': 'invoice_number', 'date facture': 'invoice_date', 'projet': 'project_name',
           'montant': 'allocation_amount', 'pourcentage': 'allocation_percent', 'taux': 'allocation_percent',
           'cout employeur': 'base_amount', 'assiette': 'base_amount', 'devise': 'currency'}

def key(value):
    import unicodedata
    return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFKD', str(value or ''))
                                      if not unicodedata.combining(c))).strip().casefold()

def decimal(value):
    text = str(value).strip().replace('\u00a0', '').replace('\u202f', '').replace(' ', '')
    text = re.sub(r'(?:EUR|€|%)$', '', text, flags=re.I).strip()
    if ',' in text:
        if text.count(',') != 1 or ('.' in text and not re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+,\d{1,2}', text)):
            raise ValueError('Nombre ambigu.')
        text = text.replace('.', '').replace(',', '.')
    if not re.fullmatch(r'-?\d+(?:\.\d+)?', text):
        raise ValueError('Nombre invalide.')
    result = Decimal(text)
    if not result.is_finite() or abs(result) > Decimal('1000000000'):
        raise ValueError('Nombre hors limites.')
    return result

def cents(value):
    number = decimal(value) * 100
    if number != number.to_integral_value():
        raise ValueError('Montant avec plus de deux décimales.')
    return int(number)

def percent(value, fmt='', mode='points'):
    number = decimal(value)
    if '%' in fmt or (mode == 'fraction' and '%' not in str(value)):
        number *= 100
    if not 0 <= number <= 100:
        raise ValueError('Pourcentage hors limites.')
    return number

def month(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m')
    value = str(value or '').strip()
    if re.fullmatch(r'\d{4}-(?:0[1-9]|1[0-2])', value):
        return value
    found = re.fullmatch(r'(0?[1-9]|1[0-2])[/.-](\d{4})', value)
    if found:
        return f'{found[2]}-{int(found[1]):02d}'
    names={'janvier':1,'fevrier':2,'mars':3,'avril':4,'mai':5,'juin':6,'juillet':7,'aout':8,'septembre':9,'octobre':10,'novembre':11,'decembre':12,'gennaio':1,'febbraio':2,'marzo':3,'aprile':4,'maggio':5,'giugno':6,'luglio':7,'agosto':8,'settembre':9,'ottobre':10,'novembre':11,'dicembre':12}
    value = key(value).replace('.', '')
    short = re.fullmatch(r'([a-z]+)[ -](\d{2}|\d{4})', value)
    if short:
        aliases = {'janv':1,'fevr':2,'fev':2,'avr':4,'juil':7,'juill':7,'sept':9,'oct':10,'nov':11,'dec':12}
        number = names.get(short[1], aliases.get(short[1]))
        if number:
            year = int(short[2]) + (2000 if len(short[2]) == 2 else 0)
            return f'{year}-{number:02d}'
    named=re.fullmatch(r'([a-z]+)\s+(\d{4})',key(value))
    if named and named[1] in names:return f'{named[2]}-{names[named[1]]:02d}'
    raise ValueError('Période attendue au format AAAA-MM.')

def read_table(raw, filename):
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError('Tableau limité à 5 Mo.')
    if filename.lower().endswith('.pdf'):
        from .stamping_tables import pdf_tables
        return pdf_tables(raw)
    if filename.lower().endswith('.csv'):
        text = raw.decode('utf-8-sig', errors='strict')
        delimiter = ';' if text.partition('\n')[0].count(';') >= text.partition('\n')[0].count(',') else ','
        data = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        if len(data) > 5001 or any(len(r) > 100 for r in data):
            raise ValueError('Limite : 5 000 lignes et 100 colonnes.')
        return {'CSV': [[{'value': v, 'raw': v, 'format': '', 'cell': f'{column(i+1)}{j+1}'} for i, v in enumerate(r)] for j, r in enumerate(data)]}
    if not filename.lower().endswith('.xlsx') or raw[:2] != b'PK':
        raise ValueError('Utilisez un fichier Excel XLSX, CSV UTF-8 ou PDF.')
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 1500 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
            raise ValueError('Classeur décompressé trop volumineux.')
        if any('vbaproject' in e.filename.lower() or e.flag_bits & 1 for e in entries):
            raise ValueError('Macros et archives chiffrées refusées.')
    from openpyxl import load_workbook
    values = load_workbook(io.BytesIO(raw), read_only=False, data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(raw), read_only=False, data_only=False, keep_links=False)
    out = {}
    if len(values.worksheets) > 25:
        raise ValueError('Maximum 25 feuilles.')
    for sheet in values:
        if sheet.max_row > 5001 or sheet.max_column > 100:
            raise ValueError('Limite : 5 000 lignes et 100 colonnes par feuille.')
        rows = []
        for row in sheet:
            cells = []
            for cell in row:
                formula = formulas[sheet.title][cell.coordinate]
                cells.append({'value': cell.value, 'raw': str(formula.value or ''), 'format': cell.number_format,
                              'cell': cell.coordinate, 'formula': formula.data_type == 'f',
                              'invalid': cell.data_type == 'e' or (formula.data_type == 'f' and cell.value is None)})
            rows.append(cells)
        # A merged value belongs only to its explicit merge, never to a neighbouring block.
        for merged in sheet.merged_cells.ranges:
            if merged.min_col == merged.max_col:
                source = rows[merged.min_row-1][merged.min_col-1]
                for r in range(merged.min_row, merged.max_row):
                    rows[r][merged.min_col-1] = {**source, 'inherited': True}
        out[sheet.title] = rows
    values.close(); formulas.close()
    return out

def column(n):
    out = ''
    while n:
        n, k = divmod(n-1, 26); out = chr(65+k) + out
    return out

def identify_header(value):
    label = key(value).replace('-', ' ').replace('’', "'")
    label = re.sub(r'\s+', ' ', label)
    label = label.replace('satalre','salaire').replace('sataire','salaire').replace('salaire.','salaire')
    if 'eligible' in label and 'projet' in label and re.search(r'montant|total|depense|cout',label):
        return 'allocation_amount', 120
    if re.search(r'(rendu(?:[ -]?compte)?|declare|impute|rapporte|retenu|eligible)', label) and re.search(r'montant|total|depense|cout',label):
        return 'allocation_amount', 100 if 'projet' in label else 80
    if ('employeur' in label and ('verse' in label or 'cout' in label)) or re.search(r'(?:salaire\s+)?brut\s+charge',label):
        return 'base_amount', 90
    if 'temps' in label and ('%' in label or 'pourcentage' in label):
        return 'allocation_percent', 90
    if re.search(r'\b(?:montant|total|cout)\b',label) and re.search(r'\b(?:ht|ttc)\b',label):
        return 'base_amount', 85 if re.search(r'\bttc\b',label) else 80
    if re.search(r'facture',label) and re.search(r'numero|n[°ºo]|reference',label):return 'invoice_number', 80
    if label in ['prestataire','nom du fournisseur','nom du prestataire','fournisseur / prestataire']:return 'supplier_ref',80
    if label in ['nom du projet','intitule du projet','projet concerne']:return 'project_name',80
    field = ALIASES.get(label, label)
    return (field, 40) if field in FIELDS else (None, 0)


def detect(rows):
    best = max(range(min(100, len(rows))), key=lambda n: sum(identify_header(c['value'])[1] for c in rows[n]), default=0)
    mapping = {}; ranks = {}
    for n,c in enumerate(rows[best] if rows else []):
        field,rank = identify_header(c['value'])
        if field and rank > ranks.get(field,0):
            mapping[field] = n; ranks[field] = rank
    return best,mapping


def complete_layout(mapping, hint, columns):
    """Carry forward headers only when at least three shared columns agree."""
    shared=set(mapping)&set(hint)
    merged={**hint,**mapping}
    if len(shared)>=3 and all(mapping[k]==hint[k] for k in shared) and len(set(merged.values()))==len(merged) and max(merged.values(),default=-1)<columns:
        return merged
    return mapping


def single_project(rows, header):
    """Use one explicit project in the preamble, never infer it from a filename."""
    found = {}
    for row in rows[:header]:
        for index, cell in enumerate(row):
            text = str(cell.get('value') or '').strip()
            names = re.findall(r'\bprojet\s*[«“"]\s*([^»”"]{1,300})[»”"]', text, re.I)
            labeled = re.fullmatch(r'(?:nom du |intitul[eé] du )?projet\s*:\s*(.{1,300})', text, re.I)
            if labeled: names.append(labeled[1])
            if key(text).rstrip(' :') in ['projet', 'nom du projet', 'intitule du projet']:
                adjacent = next((str(c.get('value') or '').strip() for c in row[index+1:] if c.get('value')), '')
                if adjacent and len(adjacent) <= 300: names.append(adjacent)
            for name in names:
                name = name.strip(); code = ''
                coded = re.fullmatch(r'((?:IF\s*)?Marittimo\s*\d+)\s*[-–—]\s*(.+)', name, re.I)
                if coded: code, name = coded[1].strip(), coded[2].strip()
                found[key(name)] = {'name':name,'code':code,'source':{'cell':cell['cell'],'raw':text}}
    return next(iter(found.values())) if len(found) == 1 else None


def inspect(raw, filename, options=None):
    from collections import Counter
    options=options or {}
    if filename.lower().endswith('.pdf') and isinstance(options.get('page'),int):
        from .stamping_tables import pdf_tables
        book=pdf_tables(raw,options['page'])
    else:book=read_table(raw,filename)
    sheets = []; results = []; unreadable = []
    required=['supplier_ref','invoice_number','project_name','allocation_amount'] if options.get('category')=='SERVICES' else ['employee_ref','period','project_name','allocation_amount']
    layouts=Counter(tuple(sorted(detect(rows)[1].items())) for rows in book.values() if all(f in detect(rows)[1] for f in required))
    common=dict(layouts.most_common(1)[0][0]) if layouts else options.get('mappingHint',{})
    for name,rows in book.items():
        best,mapping = detect(rows)
        if filename.lower().endswith('.pdf') and rows:
            mapping=complete_layout(mapping,common,len(rows[best]))
            mapping=complete_layout(mapping,options.get('mappingHint',{}),len(rows[best]))
        project = single_project(rows,best) if 'project_name' not in mapping else None
        sheets.append({'name':name,'rows':len(rows),'header':best,'mapping':mapping,
                       'preview':[[str(c['value'] if c['value'] is not None else '')[:220] for c in r] for r in rows[:max(30,best+8)]]})
        if all(f in mapping or (f=='project_name' and project) for f in required) and 'document_type' not in mapping:
            invoice=options.get('category')=='SERVICES'
            base_header=key(rows[best][mapping['base_amount']]['value']) if 'base_amount' in mapping else ''
            base_type='TTC' if re.search(r'\bttc\b',base_header) else 'HT' if re.search(r'\bht\b',base_header) else 'ELIGIBLE_BASE'
            config={'sheet':name,'header':best,'mapping':mapping,'documentType':'INVOICE' if invoice else 'PAYSLIP','baseType':base_type if invoice else 'EMPLOYER_COST','currency':'EUR','automatic':True,'grouped':not bool(project)}
            if project: config['singleProject']=project
            results.append(normalize(raw,filename,config,book))
        elif filename.lower().endswith('.pdf') and len(mapping)>=2 and ('base_amount' in mapping or 'allocation_amount' in mapping):
            unreadable.append({'row':best+1,'sheet':name,'message':'En-tête incomplet. Ajoutez la version Excel ou un PDF plus lisible pour inclure cette page.'})
    suggested = None
    if results or unreadable:
        groups = {}
        for result in results:
            for g in result['groups']:
                if g['id'] not in groups: groups[g['id']] = g
                else:
                    prior=groups[g['id']]
                    prior['allocations'].extend(g['allocations'])
                    if prior['base_amount_cents']!=g['base_amount_cents']: prior['errors'].append('Assiettes différentes entre pages.')
                    prior['errors'].extend(g['errors'])
                    update_group(prior)
        suggested={'groups':list(groups.values()),'errors':unreadable,
                   'warnings':list(dict.fromkeys(w for r in results for w in r['warnings'])),
                   'source_sha256':hashlib.sha256(raw).hexdigest(),'mapping':{'automatic':True}}
    return {'sheets':sheets,'fields':FIELDS,'sha256':hashlib.sha256(raw).hexdigest(),'suggestedImport':suggested}


def combine_inspections(parts):
    if not isinstance(parts,list) or not 1<=len(parts)<=30:raise ValueError('Pages de tableau invalides.')
    if not any(p.get('sheets') for p in parts):raise ValueError('Aucun tableau reconnu dans ce PDF. Essayez sa version Excel ou un PDF plus net.')
    first=parts[0];groups={};errors=[];warnings=[]
    for part in parts:
        if part['sha256']!=first['sha256']:raise ValueError('Les pages ne proviennent pas du même fichier.')
        result=part.get('suggestedImport')
        if not result:continue
        errors.extend(result['errors']);warnings.extend(result['warnings'])
        for g in result['groups']:
            if g['id'] not in groups:groups[g['id']]=g
            else:
                prior=groups[g['id']];prior['allocations'].extend(g['allocations']);prior['errors'].extend(g['errors'])
                if prior['base_amount_cents']!=g['base_amount_cents']:prior['errors'].append('Assiettes différentes entre pages.')
                update_group(prior)
    suggested={'groups':list(groups.values()),'errors':errors,'warnings':list(dict.fromkeys(warnings)),'source_sha256':first['sha256'],'mapping':{'automatic':True}} if groups or errors else None
    return {'sha256':first['sha256'],'fields':FIELDS,'sheets':[s for p in parts for s in p['sheets']],'suggestedImport':suggested}


def amount_value(cell):
    value = cell['value']
    if cell.get('pdf'):
        # Some PDF fonts extract the decimal comma as a space. The monetary
        # column and explicit two-digit suffix are both required to restore it.
        text = str(value).strip().replace('€','').replace('EUR','').strip()
        if re.fullmatch(r'\d{1,3}(?: \d{3})* \d{2}', text):
            text = text[:-3]+','+text[-2:]
        value=text
    if isinstance(value,(float,int)) and re.search(r'[0#]\.[0#]{2}(?![0#])',cell.get('format','')):
        return int((Decimal(str(value))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    return cents(value)


def update_group(g):
    g['total_cents']=sum(a['amount_cents'] or 0 for a in g['allocations'])
    g['remainder_cents']=None if g['base_amount_cents'] is None else g['base_amount_cents']-g['total_cents']
    g['errors']=[e for e in g['errors'] if e not in ['Total supérieur à l’assiette.','Projet présent plusieurs fois : vérifier les doublons.']]
    if g['remainder_cents'] is not None and g['remainder_cents']<0:g['errors'].append('Total supérieur à l’assiette.')
    projects=[key(a['project']) for a in g['allocations']]
    if len(projects)!=len(set(projects)):g['errors'].append('Projet présent plusieurs fois : vérifier les doublons.')


def review_import(original, changes):
    from copy import deepcopy
    if not isinstance(changes,list) or len(changes)>5000:raise ValueError('Corrections invalides.')
    result=deepcopy(original); edits={}; known={(g['id'],a['id']) for g in result['groups'] for a in g['allocations']}
    for change in changes:
        if not isinstance(change,dict) or set(change)-{'groupId','allocationId','employee','period','supplier','invoiceNumber','project','amount','baseAmount','excluded','added'}:raise ValueError('Correction invalide.')
        identity=(change.get('groupId'),change.get('allocationId'))
        if change.get('added') is True:
            g=next((g for g in result['groups'] if g['id']==change.get('groupId')),None)
            if g is None or identity in known or not re.fullmatch(r'manual-[a-f0-9-]{36}',str(change.get('allocationId',''))):raise ValueError('Nouveau projet invalide.')
            g['allocations'].append({'id':change['allocationId'],'project':'','amount_cents':0,'base_cents':g['base_amount_cents'],'programme':'','percent':None,'calculated':False,'manual':True,'source':{'rule':'MANUAL','sheet':'','cell':'','row':None}})
            known.add(identity)
        if identity not in known or identity in edits:raise ValueError('Ligne de correction inconnue ou dupliquée.')
        edits[identity]=change
    rebuilt={};errors=[];excluded=list(original.get('excluded',[]))
    for g in result['groups']:
        for a in g['allocations']:
            patch=edits.get((g['id'],a['id']),{})
            if patch.get('excluded') is True:
                excluded.append({'allocation':a,'employee':g['employee_ref'],'period':g['period']});continue
            invoice=g['document_type']=='INVOICE'
            employee=str(patch.get('supplier' if invoice else 'employee',g['supplier_ref'] if invoice else g['employee_ref'])).strip()
            invoice_number=str(patch.get('invoiceNumber',g['invoice_number'])).strip()
            period=patch.get('period',g['period']);project=str(patch.get('project',a['project'])).strip()
            value=patch.get('amount',a.get('amountRaw') if a['amount_cents'] is None else str(Decimal(a['amount_cents'])/100))
            basevalue=patch.get('baseAmount',g.get('baseRaw') if a.get('base_cents',g['base_amount_cents']) is None else str(Decimal(a.get('base_cents',g['base_amount_cents']))/100))
            try:
                period=month(period) if period or not invoice else ''
                if not 1<=len(employee)<=120 or not 1<=len(project)<=300:raise ValueError('Fournisseur et projet requis.' if invoice else 'Salarié et projet requis.')
                if invoice and not 1<=len(invoice_number)<=120:raise ValueError('Numéro de facture requis.')
                amount=cents(value)
                base=cents(basevalue) if basevalue not in [None,''] else None
                if amount<0 or (base is not None and base<=0):raise ValueError('Montant invalide.')
            except ValueError as exc:
                raise ValueError(f"{employee or 'Salarié'} · {project or 'Projet'} : {exc}") from exc
            if patch:a['review']={'before_employee':g['employee_ref'],'before_period':g['period'],'before_project':a['project'],'before_cents':a['amount_cents']}
            a['project']=project;a['amount_cents']=amount;a['base_cents']=base;a.pop('amountRaw',None)
            identity={f:g.get(f,'') for f in ['document_ref','employee_ref','invoice_number','supplier_ref','invoice_date']};identity['supplier_ref' if invoice else 'employee_ref']=employee
            if invoice:identity['invoice_number']=invoice_number
            gid=hashlib.sha256(json.dumps([g['document_type'],period,identity],sort_keys=True).encode()).hexdigest()[:24]
            target=rebuilt.setdefault(gid,{**g,**identity,'id':gid,'period':period,'base_amount_cents':base,'allocations':[],'errors':[],'rounding_cents':None})
            if target['base_amount_cents']!=base:target['errors'].append('Assiettes différentes dans un même groupe.')
            target['allocations'].append(a)
    for g in rebuilt.values():update_group(g)
    result.update(groups=list(rebuilt.values()),errors=list(original.get('errors',[]))+errors,excluded=excluded)
    return result


def normalize(raw, filename, config, book=None):
    book = book if book is not None else read_table(raw, filename)
    sheet = config.get('sheet', next(iter(book)))
    if sheet not in book: raise ValueError('Feuille introuvable.')
    rows = book[sheet]; mapping = config.get('mapping', {})
    header = config.get('header', 0)
    if not isinstance(header, int) or not 0 <= header < len(rows): raise ValueError('En-tête invalide.')
    if any(k not in FIELDS or not isinstance(v, int) or not 0 <= v < 100 for k,v in mapping.items()):
        raise ValueError('Correspondance invalide.')
    groups = {}; errors = []; warnings = []; pending = []; prior = {}; prior_cells = {}; source_hash = hashlib.sha256(raw).hexdigest()
    excluded = config.get('excludedRows', [])
    if excluded and not str(config.get('exclusionReason', '')).strip(): raise ValueError('Justifiez les lignes exclues.')
    wide = config.get('projects', [])
    if len(wide) > 60: raise ValueError('Maximum 60 colonnes projet.')
    for index, cells in enumerate(rows[header+1:], header+2):
        if index in excluded: continue
        if not any(c['value'] is not None and str(c['value']).strip() for c in cells):
            prior = {}; prior_cells = {}; continue
        def cell(field):
            n = mapping.get(field)
            return cells[n] if n is not None and n < len(cells) else {'value': None, 'format': '', 'cell': '', 'raw': ''}
        def present(value):
            return None if isinstance(value,str) and not value.strip() else value
        data = {f: present(cell(f)['value']) for f in FIELDS}
        source_cells = {f:cell(f) for f in mapping}
        if config.get('singleProject'):
            # Empty payroll rows and the adjacent charges ledger are not allocations.
            # Keep incomplete *populated* monetary rows for explicit correction.
            if data['allocation_amount'] is None and data['base_amount'] is None: continue
            if key(data['employee_ref']).startswith('total') or key(data['period']).startswith('total'): continue
            project = config['singleProject']
            data['project_name'] = data['project_name'] or project['name']
            data['project_code'] = data['project_code'] or project.get('code','')
            source_cells['project_name'] = {**project['source'], 'value':project['name']}
        if config.get('grouped'):
            identity_fields = ['document_type','document_ref','period','employee_ref','invoice_number','invoice_date','supplier_ref','base_amount','base_type','currency']
            # Any new identifier starts a block. No inheritance across an empty separator.
            changed = any(data[f] is not None and data[f] != prior.get(f) for f in ['document_ref','period','employee_ref','invoice_number','supplier_ref'])
            if changed: prior = {}; prior_cells = {}
            for f in identity_fields:
                if data[f] is not None: prior[f] = data[f]; prior_cells[f] = cell(f)
                else:
                    data[f] = prior.get(f)
                    if f in prior_cells: source_cells[f] = prior_cells[f]
        if config.get('automatic'):
            project_text=key(data['project_name'])
            if not project_text or project_text in ['projet','projets','total'] or project_text.startswith('total '): continue
        variants = wide or [None]
        for project in variants:
            try:
                current = dict(data); amount_cell = cell('allocation_amount'); rate_cell = cell('allocation_percent')
                if project:
                    current['project_name'] = project['name']
                    n = project.get('amountColumn'); p = project.get('percentColumn')
                    amount_cell = cells[n] if isinstance(n,int) and 0 <= n < len(cells) else {'value':None,'cell':'','raw':''}
                    rate_cell = cells[p] if isinstance(p,int) and 0 <= p < len(cells) else {'value':None,'cell':'','raw':'','format':''}
                    current['allocation_amount'] = present(amount_cell['value']); current['allocation_percent'] = present(rate_cell['value'])
                    if current['allocation_amount'] is None and current['allocation_percent'] is None: continue
                # Group title rows are not allocations. They may establish the explicitly selected block.
                if not current['project_name'] and config.get('grouped') and current['allocation_amount'] is None: continue
                used = list(source_cells.values()) + [amount_cell, rate_cell]
                if any(c.get('invalid') for c in used): raise ValueError('Cellule en erreur ou formule sans valeur enregistrée.')
                if any(c.get('formula') for c in used): warnings.append('Formules : valeurs enregistrées utilisées, à confirmer dans Excel.')
                dtype = str(current['document_type'] or config.get('documentType') or '').upper()
                dtype = {'PAIE':'PAYSLIP','BULLETIN':'PAYSLIP','FACTURE':'INVOICE'}.get(dtype, dtype)
                if dtype not in ['PAYSLIP','INVOICE']: raise ValueError('Type de pièce requis.')
                period = month(current['period']) if current['period'] else ''
                if dtype == 'PAYSLIP' and not period: raise ValueError('Mois de paie requis.')
                name = str(current['project_name'] or '').strip()
                if not name or len(name) > 300: raise ValueError('Nom de projet requis (300 caractères maximum).')
                currency = str(current['currency'] or config.get('currency') or '').upper()
                if currency != 'EUR': raise ValueError('Devise EUR à confirmer. Autres devises non prises en charge.')
                base = amount_value({**cell('base_amount'),'value':current['base_amount']}) if current['base_amount'] is not None and current['base_amount'] != '' else None
                base_type = str(current['base_type'] or ('EMPLOYER_COST' if dtype=='PAYSLIP' else config.get('baseType',''))).upper()
                if base_type not in ['EMPLOYER_COST','HT','TTC','ELIGIBLE_BASE']: raise ValueError('Assiette à préciser (HT, TTC, coût employeur, éligible).')
                if dtype=='PAYSLIP' and base_type!='EMPLOYER_COST': raise ValueError('La paie doit utiliser le coût employeur.')
                try:
                    rate = percent(re.sub(r'(?<=\d) (\d{2})%',r',\1%',str(current['allocation_percent'])) if rate_cell.get('pdf') else current['allocation_percent'],rate_cell.get('format',''),config.get('percentMode','points')) if current['allocation_percent'] not in [None,''] else None
                except ValueError:
                    if not config.get('automatic') or current['allocation_amount'] in [None,'']:raise
                    rate=None;warnings.append('Certains pourcentages sont illisibles. Les montants explicites du tableau sont conservés.')
                calculated = current['allocation_amount'] in [None,'']
                if calculated:
                    if not config.get('allowCalculation') or rate is None or base is None: raise ValueError('Montant absent : calcul au pourcentage non autorisé ou base absente.')
                    amount = int((Decimal(base)*rate/100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
                else: amount = amount_value(amount_cell)
                if amount < 0 or (base is not None and base <= 0): raise ValueError('Montants négatifs et assiettes nulles non autorisés.')
                identity = {f:str(current[f] or '').strip() for f in ['document_ref','employee_ref','invoice_number','supplier_ref','invoice_date']}
                if isinstance(current['invoice_date'],(datetime,date)): identity['invoice_date']=current['invoice_date'].strftime('%Y-%m-%d')
                if identity['invoice_date']:
                    if re.fullmatch(r'\d{2}/\d{2}/\d{4}',identity['invoice_date']):identity['invoice_date']=datetime.strptime(identity['invoice_date'],'%d/%m/%Y').date().isoformat()
                    date.fromisoformat(identity['invoice_date'])
                if dtype=='PAYSLIP' and not identity['employee_ref']: raise ValueError('Matricule ou référence salarié requis.')
                if dtype=='INVOICE' and (not identity['invoice_number'] or not identity['supplier_ref']): raise ValueError('Numéro facture et référence fournisseur requis.')
                gid = hashlib.sha256(json.dumps([dtype,period,identity],sort_keys=True).encode()).hexdigest()[:24]
                g = groups.setdefault(gid, {'id':gid, 'document_type':dtype,'period':period,**identity,'currency':currency,'base_amount_cents':base,'base_type':base_type,'allocations':[],'errors':[]})
                if g['base_amount_cents']!=base or g['base_type']!=base_type: g['errors'].append('Assiettes différentes dans un même groupe.')
                if any(key(a['project'])==key(name) for a in g['allocations']): g['errors'].append('Projet présent plusieurs fois : vérifier les doublons.')
                g['allocations'].append({'id':f'{hashlib.sha256(sheet.encode()).hexdigest()[:8]}-{index}-{len(g["allocations"])}','project':name,'code':str(current['project_code'] or ''),'programme':str(current['programme'] or ''),'amount_cents':amount,'base_cents':base,'percent':str(rate) if rate is not None else None,'calculated':calculated,
                  'source':{'file':filename,'sha256':source_hash,'sheet':sheet,'row':index,'cell':amount_cell['cell'] if not calculated else rate_cell['cell'],'raw':str(amount_cell.get('raw','')),'rule':'ROUND_HALF_UP' if calculated else ('DISPLAYED_CENTS' if amount_cell.get('formula') else 'EXACT'),'fields':{f:{'cell':c['cell'],'raw':str(c['raw'])} for f,c in source_cells.items()}}})
            except (ValueError,InvalidOperation,KeyError,TypeError) as exc:
                message=str(exc) or 'Ligne invalide.'
                errors.append({'row':index,'sheet':sheet,'message':message})
                if config.get('automatic'):
                    try: pending_amount=amount_value(amount_cell)
                    except (ValueError,TypeError,InvalidOperation):pending_amount=None
                    try: pending_base=amount_value({**cell('base_amount'),'value':current['base_amount']}) if current['base_amount'] not in [None,''] else None
                    except (ValueError,TypeError,InvalidOperation):pending_base=None
                    try: pending_period=month(current['period'])
                    except ValueError:pending_period=str(current['period'] or '')
                    aid=f'{hashlib.sha256(sheet.encode()).hexdigest()[:8]}-{index}-pending'
                    pending.append({'id':aid,'document_type':config.get('documentType','PAYSLIP'),'employee_ref':str(current['employee_ref'] or ''),'period':pending_period,'document_ref':'','invoice_number':str(current['invoice_number'] or ''),'supplier_ref':str(current['supplier_ref'] or ''),'invoice_date':str(current['invoice_date'] or ''),'currency':'EUR','base_type':config.get('baseType','EMPLOYER_COST'),'base_amount_cents':pending_base,'baseRaw':str(current['base_amount'] or ''),'errors':[message],'total_cents':pending_amount or 0,'remainder_cents':None,'rounding_cents':None,
                        'allocations':[{'id':aid,'project':str(current['project_name'] or ''),'amount_cents':pending_amount,'amountRaw':str(current['allocation_amount'] or ''),'code':'','programme':'','percent':None,'calculated':False,'source':{'file':filename,'sha256':source_hash,'sheet':sheet,'row':index,'cell':amount_cell['cell'],'raw':str(amount_cell.get('raw','')),'rule':'REVIEW_REQUIRED'}}]})
    for group in groups.values():
        group['total_cents'] = sum(a['amount_cents'] for a in group['allocations'])
        group['remainder_cents'] = None if group['base_amount_cents'] is None else group['base_amount_cents']-group['total_cents']
        if group['remainder_cents'] is not None and group['remainder_cents']<0: group['errors'].append('Total supérieur à l’assiette.')
        rates=[a['percent'] for a in group['allocations']]
        group['rounding_cents']=None
        if all(r is not None for r in rates) and group['base_amount_cents'] is not None:
            rate=sum(Decimal(r) for r in rates)
            if rate>100: group['errors'].append('Pourcentages supérieurs à 100 %.')
            group['rounding_cents']=group['total_cents']-int((Decimal(group['base_amount_cents'])*rate/100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    return {'groups':list(groups.values())+pending,'errors':errors,'warnings':list(dict.fromkeys(warnings)), 'source_sha256':source_hash,'mapping':config}
